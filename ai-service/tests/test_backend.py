import os
import unittest
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import patch

os.environ.update(DB_URL="postgresql+psycopg://test:test@localhost/test",
                  AI_SERVICE_TOKEN="a"*32,OPENROUTER_API_KEY="test",
                  MINIO_ACCESS_KEY="test",MINIO_SECRET_KEY="test")

from fastapi.testclient import TestClient
from pydantic import ValidationError
from app import main, examiner, ingest
from app.schemas import GradeAttemptRequest, GeneratedExam, GradeResult
from app.extractors import extract_text

TOKEN={"X-Service-Token":"a"*32}
client=TestClient(main.app)

def response(content):
    return SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(content=content))])

class Contracts(unittest.TestCase):
    def test_auth_required_on_all_internal_routes(self):
        for endpoint,body in [("/index",{"object_key":"x"}),
                              ("/ask",{"question":"q","lesson_id":1}),
                              ("/exams/generate",{"lesson_id":1}),
                              ("/exams/grade",{"questions":[{"id":1,"text":"q","max_score":10}],"answers":[]})]:
            with self.subTest(endpoint=endpoint):
                self.assertEqual(client.post(endpoint,json=body).status_code,401)
                self.assertEqual(client.post(endpoint,json=body,headers={"X-Service-Token":"wrong"}).status_code,401)

    def test_invalid_count(self):
        for count in [0,21]:
            self.assertEqual(client.post("/exams/generate",json={"lesson_id":1,"count":count},headers=TOKEN).status_code,422)

    def test_blank_question(self):
        self.assertEqual(client.post("/ask",json={"lesson_id":1,"question":"   "},headers=TOKEN).status_code,422)

    def test_java_camel_and_python_snake_aliases(self):
        for score_key,id_key,text_key in [("maxScore","questionId","answerText"),("max_score","question_id","answer_text")]:
            body={"questions":[{"id":1,"text":"q",score_key:10,"chunks":["context"]}],
                  "answers":[{id_key:1,text_key:"answer"}]}
            self.assertEqual(GradeAttemptRequest.model_validate(body).answers[0].question_id,1)

    def test_duplicate_question_rejected(self):
        q={"id":1,"text":"q","max_score":10}
        with self.assertRaises(ValidationError): GradeAttemptRequest(questions=[q,q],answers=[])

    def test_duplicate_answer_rejected(self):
        a={"question_id":1,"answer_text":"a"}
        with self.assertRaises(ValidationError):
            GradeAttemptRequest(questions=[{"id":1,"text":"q","max_score":10}],answers=[a,a])

    def test_unknown_question_rejected(self):
        with self.assertRaises(ValidationError):
            GradeAttemptRequest(questions=[{"id":1,"text":"q","max_score":10}],
                                answers=[{"question_id":2,"answer_text":"a"}])

    def test_missing_answer_graded(self):
        body={"questions":[{"id":1,"text":"q","maxScore":10}],"answers":[]}
        result=client.post("/exams/grade",json=body,headers=TOKEN)
        self.assertEqual(result.status_code,200)
        self.assertEqual(result.json()[0]["score"],0)

    def test_ai_invalid_output_is_not_client_error_or_details_leak(self):
        with patch.object(main,"generate_exam",side_effect=ValidationError.from_exception_data("AI",[])):
            result=client.post("/exams/generate",json={"lesson_id":1},headers=TOKEN)
        self.assertEqual(result.status_code,502)

    def test_health_uses_existing_model(self):
        with patch.object(main.engine,"connect") as connection:
            connection.return_value.__enter__.return_value.execute.return_value=None
            result=client.get("/health")
        self.assertEqual(result.status_code,200)
        self.assertIn("llm_model",result.json())

class Processing(unittest.TestCase):
    def test_short_material_is_kept(self):
        self.assertEqual(ingest.chunk_text("short text"),["short text"])

    def test_chunk_overlap_cannot_loop(self):
        for size,overlap in [(0,0),(5,5),(5,6),(5,-1)]:
            with self.assertRaises(ValueError): ingest.chunk_text("text",size,overlap)

    def test_chunks_cover_text_without_extra_tail(self):
        self.assertEqual(ingest.chunk_text("abcdefghij",6,2),["abcdef","efghij"])

    def test_utf8_and_corrupt_text(self):
        self.assertEqual(extract_text("Привет".encode(),"text/plain","a.txt"),"Привет")
        with self.assertRaises(ValueError): extract_text(b"\xff",None,"a.txt")

    def test_docx_tables(self):
        from docx import Document
        doc=Document(); doc.add_paragraph("Lesson"); doc.add_table(rows=1,cols=1).cell(0,0).text="Table"
        out=BytesIO(); doc.save(out)
        self.assertIn("Table",extract_text(out.getvalue(),None,"a.docx"))

    def test_corrupt_pdf(self):
        with self.assertRaises(ValueError): extract_text(b"not pdf",None,"a.pdf")

    def test_unsupported_format(self):
        with self.assertRaises(ValueError): extract_text(b"text",None,"a.exe")

    def test_embedding_count_mismatch(self):
        with patch.object(ingest.httpx,"Client") as http:
            http.return_value.__enter__.return_value.post.return_value.json.return_value={"embeddings":[]}
            with self.assertRaises(ValueError): ingest.embed(["text"])

    def test_embedding_dimension_mismatch(self):
        with patch.object(ingest.httpx,"Client") as http:
            http.return_value.__enter__.return_value.post.return_value.json.return_value={"embeddings":[[1]]}
            with self.assertRaises(ValueError): ingest.embed(["text"])

    def test_bad_vectors_do_not_delete_old_index(self):
        with patch.object(ingest,"embed",return_value=[]),patch.object(ingest.engine,"begin") as begin:
            with self.assertRaises(ValueError): ingest.index_material(1,b"text",None,"a.txt")
            begin.assert_not_called()

    def test_generation_count_exact(self):
        with patch.object(examiner,"get_random_chunks",return_value=["lesson"]), \
             patch.object(examiner.client.chat.completions,"create",return_value=response('{"questions":[{"question":"q","max_score":10}]}')):
            with self.assertRaises(ValueError): examiner.generate_exam(1,2)

    def test_generated_score_is_strict(self):
        for score in [0,11,"10",True]:
            with self.assertRaises(ValidationError):
                GeneratedExam.model_validate({"questions":[{"question":"q","max_score":score}]})

    def test_grade_does_not_accept_score_above_max(self):
        with patch.object(examiner.client.chat.completions,"create",return_value=response('{"score":11,"feedback":"x"}')):
            with self.assertRaises(ValueError):
                examiner.grade_attempt([{"question_id":1,"question":"q","max_score":10,"chunks":["context"],"answer_text":"answer"}])

if __name__=="__main__": unittest.main()
